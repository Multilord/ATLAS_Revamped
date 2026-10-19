"""Historical import, local evidence search and ticket references; no AI calls."""
import base64
import csv
import io
import re
import zipfile
from pathlib import Path
from datetime import datetime, timezone
import openpyxl
from ingestion import MAX_FILE, MAX_TEXT, extract_upload, prepare_source, chunks, now
from storage import ConflictError
import security

KINDS = ('Historical ticket', 'Resolution guide', 'Investigation report', 'Known issue', 'Change record')
QUALITIES = ('Unreviewed', 'Reviewed', 'Verified fix', 'Workaround', 'Unresolved findings')
DEMO_CUSTOMER = 'ATLAS Historical Demo'
FIELDS = {'title': ('title', 'summary', 'short_description'),
    'content': ('problem_statement', 'description', 'content'),
    'source_reference': ('ticket_id', 'number', 'id'),
    'application': ('application', 'module'), 'platform': ('platform',), 'module': ('module',),
    'version': ('release', 'version'), 'component': ('component',),
    'solution': ('resolution', 'solution'), 'source_updated_at': ('closed', 'updated_at', 'updated')}


def text(data, key, required=False, limit=12000):
    value = data.get(key, '')
    if not isinstance(value, str) or len(value) > limit or (required and not value.strip()):
        raise ValueError('Invalid or missing ' + key.replace('_', ' '))
    return value.strip()


def synthetic(record):
    return record.get('synthetic') is True or record.get('demo') is True or record.get('data_origin') == 'Synthetic demo'


def permitted(source, ticket):
    scoped = source.get('customer') == ticket.get('customer') or (source.get('customer') == 'General' and source.get('generalization_confirmed') is True)
    demo_library = source.get('historical_library') == 'synthetic-v1' and synthetic(ticket)
    return (scoped or demo_library) and (not synthetic(source) or synthetic(ticket))


def quality(source):
    if source.get('evidence_quality') == 'Unresolved findings':
        return 'Unresolved findings'
    if source.get('solution_kind') == 'Workaround':
        return 'Workaround'
    if source.get('validation') == 'Verified':
        return 'Verified fix'
    return 'Reviewed' if source.get('validation') in ('Reviewed', 'Outdated') else 'Unreviewed'


def preview(data):
    filename = text(data, 'filename', True, 255)
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.csv', '.xlsx', '.xlsm', '.xltx'):
        extracted = extract_upload(data)
        return dict(mode='document', filename=extracted['filename'], content=extracted['content'], sheets=[])
    try:
        raw = base64.b64decode(data.get('file_base64', ''), validate=True)
    except (ValueError, TypeError):
        raise ValueError('Invalid file encoding.') from None
    if not raw or len(raw) > MAX_FILE:
        raise ValueError('Upload a non-empty file of at most 5 MB.')
    if suffix == '.csv':
        try:
            matrix = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
        except (UnicodeError, csv.Error):
            raise ValueError('Use a UTF-8 CSV with a header row.') from None
        sheets = []
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(i.file_size for i in members) > 20 * 1024 * 1024 or any(i.file_size > 10 * 1024 * 1024 or (i.file_size > 100000 and i.file_size / max(1, i.compress_size) > 200) for i in members):
                    raise ValueError('Spreadsheet expands beyond the permitted size.')
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                sheets = book.sheetnames
                selected = data.get('sheet') or sheets[0]
                if selected not in sheets:
                    raise ValueError('Select an existing worksheet.')
                matrix = []
                for row in book[selected].iter_rows(values_only=True):
                    if any(v is not None and str(v).strip() for v in row):
                        matrix.append([v.isoformat() if isinstance(v, datetime) else str(v) if v is not None else '' for v in row])
                    if len(matrix) > 101:
                        raise ValueError('Import at most 100 records at a time.')
            finally:
                book.close()
        except (ValueError,):
            raise
        except Exception:
            raise ValueError('Cannot read this workbook. Use a valid XLSX file.') from None
    matrix = [r for r in matrix if any(str(c).strip() for c in r)]
    if not 2 <= len(matrix) <= 101:
        raise ValueError('Include one header row and 1–100 data rows.')
    headers = [str(c).strip() for c in matrix[0]]
    if len(headers) > 60 or any(not h or len(h) > 150 for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('Use unique, non-empty column names (at most 60 columns).')
    rows = []
    for i, row in enumerate(matrix[1:], 2):
        if len(row) != len(headers):
            raise ValueError(f'Row {i} has a different number of columns.')
        rows.append(dict(zip(headers, map(str, row))))
    if sum(len(v) for r in rows for v in r.values()) > MAX_TEXT:
        raise ValueError('Import text exceeds 120,000 characters. Split the file.')
    mapping = {field: next((h for alias in aliases for h in headers if h.casefold().replace(' ', '_') == alias), '') for field, aliases in FIELDS.items()}
    return dict(mode='table', filename=Path(filename).name, headers=headers, rows=rows, mapping=mapping,
        sheets=sheets, sheet=data.get('sheet') or (sheets[0] if sheets else ''))


def parse_date(value, date_format):
    if not value:
        return ''
    try:
        parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
    except ValueError:
        if date_format not in ('MM/DD/YYYY', 'DD/MM/YYYY'):
            raise ValueError('Select the date format for non-ISO source dates.') from None
        try:
            parsed = datetime.strptime(value, '%m/%d/%Y' if date_format == 'MM/DD/YYYY' else '%d/%m/%Y')
        except ValueError:
            raise ValueError('A source date does not match the selected format.') from None
    return parsed.replace(tzinfo=parsed.tzinfo or timezone.utc).isoformat()


def prepare_import(data, store):
    defaults = data.get('defaults', {})
    if not isinstance(defaults, dict):
        raise ValueError('Invalid import metadata.')
    filename = text(data, 'filename', False, 255)
    rows = data.get('rows')
    candidates = []
    if rows is not None:
        mapping = data.get('mapping', {})
        if not isinstance(rows, list) or not 1 <= len(rows) <= 100 or not isinstance(mapping, dict):
            raise ValueError('Import 1–100 rows with a valid field mapping.')
        if any(not isinstance(r, dict) or len(r) > 60 or any(not isinstance(k, str) or not isinstance(v, str) for k, v in r.items()) for r in rows):
            raise ValueError('Invalid table rows.')
        if sum(len(v) for r in rows for v in r.values()) > MAX_TEXT:
            raise ValueError('Import text exceeds 120,000 characters.')
        if not mapping.get('title') or not mapping.get('content'):
            raise ValueError('Map the title and problem/content columns.')
        if any(k not in FIELDS or not isinstance(v, str) or (v and any(v not in r for r in rows)) for k,v in mapping.items()):
            raise ValueError('Invalid field mapping.')
        for index, row in enumerate(rows, 2):
            values = {k: row[column].strip() for k,column in mapping.items() if column}
            # Preserve the complete row, including unsuccessful attempts and verification notes.
            content = '\n'.join(f'{k}: {v}' for k,v in row.items() if v.strip())
            if not values.get('title') or not values.get('content'):
                raise ValueError(f'Row {index} is missing title or problem/content.')
            values['content'] = content
            values['source_updated_at'] = parse_date(values.get('source_updated_at', ''), data.get('date_format', 'ISO'))
            candidates.append(({**defaults, **values, 'filename': filename}, index))
    else:
        candidates.append(({**defaults, 'content': text(data, 'content', True, MAX_TEXT), 'filename': filename}, None))
    documents = []
    for values, row_number in candidates:
        kind = values.get('source_type', 'Historical ticket')
        if kind not in KINDS:
            raise ValueError('Select a historical document classification.')
        evidence = values.get('evidence_quality', 'Unreviewed')
        if evidence not in QUALITIES:
            raise ValueError('Select a valid evidence quality.')
        is_demo = values.get('synthetic', False)
        if not isinstance(is_demo, bool):
            raise ValueError('Synthetic flag must be true or false.')
        if values.get('customer') == 'General' and values.get('generalized_confirmed') is not True:
            raise ValueError('Confirm de-identification before sharing General knowledge.')
        if is_demo and values.get('customer') == 'General':
            raise ValueError('Synthetic sources must use a dedicated customer scope, not General.')
        values['source_type'] = kind
        values['solution_kind'] = 'Workaround' if evidence == 'Workaround' else 'Unknown' if evidence == 'Unresolved findings' else 'Likely solution'
        doc = prepare_source(values, store)
        doc.update(historical=True, component=text(values, 'component', limit=150),
            document_version=text(values, 'document_version', limit=100),
            evidence_quality=evidence, synthetic=is_demo,
            data_origin='Synthetic demo' if is_demo else 'User upload',
            import_row=row_number, imported_by=security.actor() if security.principal.get() else 'Consultant')
        if evidence in ('Reviewed', 'Verified fix'):
            doc['review_notes'] = text(values, 'review_notes', True)
            doc.update(validation='Reviewed', reviewed_at=now())
        if evidence == 'Verified fix':
            checks = {k: values.get(k, False) for k in ('tested', 'customer_confirmed', 'deployed')}
            if any(v is not True for v in checks.values()):
                raise ValueError('Verified fixes require test, customer confirmation and deployment evidence.')
            if is_demo:
                doc.update(simulated_quality='Verified fix', evidence_quality='Reviewed')
            else:
                doc.update(validation='Verified', validation_evidence=checks, solution_kind='Verified fix')
        documents.append(doc)
    return documents


def import_documents(data, store):
    docs = prepare_import(data, store)  # Validate every row before any writes.
    saved, duplicates = [], []
    for doc in docs:
        existing = store.db.sources.find_one({'customer': doc['customer'], 'content_hash': doc['content_hash']}, {'_id': 0})
        if existing:
            duplicates.append(existing['id'])
            continue
        try:
            store.save(doc, 'source')
            saved.append(doc['id'])
        except ConflictError:
            existing = store.db.sources.find_one({'customer': doc['customer'], 'content_hash': doc['content_hash']})
            if not existing:
                raise
            duplicates.append(existing['id'])
    return dict(saved=saved, duplicates=duplicates, indexing_requests=0)


def index_estimate(source, model):
    count = len(chunks(source['content']))
    return dict(chunks=count, requests=(count+15)//16, already_indexed=source.get('index_status') == 'Ready' and source.get('embedding_model') == model)


def search(ticket, data, store):
    terms = set(re.findall(r'[\w-]{2,}', text(data, 'query', True, 500).casefold()))
    found = []
    for source in store.records('source'):
        if not permitted(source, ticket):
            continue
        haystack = (source['title'] + ' ' + source['content']).casefold()
        matches = sum(term in haystack for term in terms)
        if not matches:
            continue
        passage = max(chunks(source['content']), key=lambda c: sum(term in c['text'].casefold() for term in terms))
        found.append(dict(source_id=source['id'], title=source['title'], excerpt=passage['text'],
            section=source.get('source_reference') or source.get('filename', ''), evidence_quality=quality(source),
            synthetic=synthetic(source), validation=source['validation'], matches=matches,
            version=source.get('version', ''), source_updated_at=source.get('source_updated_at', '')))
    return dict(results=sorted(found, key=lambda r: (-r['matches'], r['title']))[:20])


def attach(ticket, data, store, author):
    if data.get('revision') != ticket.get('revision', 0):
        raise ConflictError('Ticket changed. Reload before attaching evidence.')
    if ticket.get('lifecycle_stage') == 'Closure':
        raise ValueError('The ticket lifecycle is closed. Attach evidence to an open investigation.')
    source_id = text(data, 'source_id', True, 100)
    source = store.get(source_id, 'source')
    if not source or not permitted(source, ticket):
        raise ValueError('Source is not permitted for this ticket or its demo scope.')
    quote = text(data, 'quote', False, 3000)
    if quote and ' '.join(quote.split()).casefold() not in ' '.join(source['content'].split()).casefold():
        raise ValueError('The quotation must match the original source text.')
    reason = text(data, 'reason', True, 2000)
    links = ticket.setdefault('source_links', [])
    if len(links) >= 100:
        raise ValueError('At most 100 evidence references per ticket.')
    if any(r['source_id'] == source_id and r.get('quote', '') == quote for r in links):
        raise ValueError('This reference is already attached.')
    stamp = now()
    links.append(dict(source_id=source_id, title=source['title'], quote=quote, reason=reason,
        evidence_quality=quality(source), synthetic=synthetic(source), content_hash=source['content_hash'],
        attached_at=stamp, attached_by=author))
    ticket.setdefault('timeline', []).append(dict(at=stamp, kind='Attachment',
        text=f'Historical reference: {source["title"]}. {reason}' + (f'\nQuoted evidence: {quote}' if quote else ''), author=author, source=source_id))
    # A reference is not an accepted recommendation and does not advance lifecycle.
    ticket['analysis'] = None
    ticket['updated_at'] = stamp
    store.save(ticket)
    return ticket
