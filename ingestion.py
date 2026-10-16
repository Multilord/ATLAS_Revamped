"""Bounded document extraction and atomic indexing; no fabricated records."""
import base64
import binascii
import hashlib
from html.parser import HTMLParser
import io
import json
import re
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from pypdf import PdfReader
from docx import Document
import openpyxl
from ai import AIError

MAX_FILE = 5 * 1024 * 1024
MAX_TEXT = 120000


class _VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'noscript'):
            self.hidden += 1
        elif tag in ('p', 'div', 'br', 'li', 'tr', 'h1', 'h2', 'h3', 'section'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'noscript') and self.hidden:
            self.hidden -= 1
        elif tag in ('p', 'div', 'li', 'tr', 'h1', 'h2', 'h3', 'section'):
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.parts.append(data.strip())


def now():
    return datetime.now(timezone.utc).isoformat()


def extract_upload(data):
    filename = data.get('filename', '')
    if not isinstance(filename, str) or len(filename) > 255:
        raise ValueError('Invalid filename.')
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.txt', '.md', '.json', '.pdf', '.docx', '.pptx', '.html', '.htm', '.xlsx', '.xlsm', '.xltx', '.csv'):
        raise ValueError('Upload TXT, MD, JSON, PDF, DOCX, PPTX, HTML, XLSX or CSV files.')
    try:
        content = base64.b64decode(data.get('file_base64', ''), validate=True)
    except (ValueError, TypeError, binascii.Error):
        raise ValueError('Invalid file encoding.') from None
    if not content or len(content) > MAX_FILE:
        raise ValueError('Files must be non-empty and at most 5 MB.')
    try:
        if suffix == '.pdf':
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted or len(reader.pages) > 150:
                raise ValueError('Use an unencrypted PDF with at most 150 pages.')
            parts = []
            for index, page in enumerate(reader.pages):
                parts.append(f'[Page {index+1}]\n' + (page.extract_text() or ''))
                if sum(len(p) for p in parts) > MAX_TEXT:
                    raise ValueError('Extracted text exceeds 120,000 characters. Split the document.')
            text = '\n\n'.join(parts)
            if not any(re.search(r'\w{3}', p.split('\n', 1)[-1]) for p in parts):
                raise ValueError('No readable PDF text found. Run OCR on scanned documents before upload.')
        elif suffix == '.docx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('DOCX expands beyond the permitted size.')
            doc = Document(io.BytesIO(content))
            text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        elif suffix == '.pptx':
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if sum(i.file_size for i in members) > 20 * 1024 * 1024 or len(members) > 2000:
                    raise ValueError('PowerPoint expands beyond the permitted size.')
                slide_names = sorted((i.filename for i in members if re.fullmatch(r'ppt/slides/slide\d+\.xml', i.filename)),
                    key=lambda name: int(re.search(r'slide(\d+)', name).group(1)))
                if len(slide_names) > 200:
                    raise ValueError('Use a PowerPoint with at most 200 slides.')
                parts = []
                for slide_number, name in enumerate(slide_names, 1):
                    root = ET.fromstring(archive.read(name))
                    lines = [element.text or '' for element in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t')]
                    if lines:
                        parts.append(f'[Slide: {slide_number}]\n' + ' '.join(lines))
                text = '\n\n'.join(parts)
        elif suffix in ('.html', '.htm'):
            parser = _VisibleHTML()
            parser.feed(content.decode('utf-8-sig', errors='replace'))
            parser.close()
            text = '\n'.join(part for part in parser.parts if part.strip())
        elif suffix in ('.xlsx', '.xlsm', '.xltx'):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Spreadsheet expands beyond the permitted size.')
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            try:
                if len(wb.sheetnames) > 100:
                    raise ValueError('Use a spreadsheet with at most 100 worksheets.')
                parts, extracted = [], 0
                for sheetname in wb.sheetnames:
                    rows = []
                    for row in wb[sheetname].iter_rows(values_only=True):
                        if not any(c is not None and str(c).strip() for c in row):
                            continue
                        line = ' | '.join(str(c).strip() if c is not None else '' for c in row)
                        extracted += len(line) + 1
                        if extracted > MAX_TEXT:
                            raise ValueError('Extracted text exceeds 120,000 characters. Split the spreadsheet.')
                        rows.append(line)
                    if rows:
                        parts.append(f'[Sheet: {sheetname}]\n' + '\n'.join(rows))
                text = '\n\n'.join(parts)
            finally:
                wb.close()
        elif suffix == '.csv':
            text = content.decode('utf-8-sig', errors='replace')
        else:
            text = content.decode('utf-8-sig')
            if suffix == '.json':
                json.loads(text)
    except ValueError:
        raise
    except Exception:
        raise ValueError('The file could not be read. Check its format and try again.') from None
    if not text.strip() or len(text) > MAX_TEXT:
        raise ValueError('Extracted text must contain 1–120,000 characters. Split larger documents.')
    return {'content': text, 'filename': Path(filename).name}


def chunks(text):
    output, start = [], 0
    while start < len(text):
        end = min(start + 2200, len(text))
        if end < len(text):
            boundary = text.rfind('\n', start + 1600, end)
            if boundary != -1:
                end = boundary
        piece = text[start:end].strip()
        if piece:
            output.append({'id': f'C{len(output)+1:04}', 'text': piece, 'offset': start})
        if end == len(text):
            break
        start = end - 200
    return output


def index_source(source, store, ai):
    if source['customer'] != store.customer and store.customer not in ('My Workspace', 'All') and source['customer'] != 'General':
        raise ValueError('This source is outside the writable workspace.')
    pieces = chunks(source['content'])
    try:
        vectors = ai.embed([source['title'] + '\n' + part['text'] for part in pieces])
        if len(vectors) != len(pieces) or len({len(v) for v in vectors}) != 1:
            raise AIError('Incomplete embedding response. Retry indexing.')
        for part, vector in zip(pieces, vectors):
            part['embedding'] = vector
        source.update(chunks=pieces, chunk_count=len(pieces), embedding_model=ai.config.embedding_model,
            index_status='Ready', index_error='', indexed_at=now())
    except AIError as exc:
        source.update(chunks=[], chunk_count=0, index_status='Error', index_error=str(exc))
    store.save(source, 'source')
    return {k: v for k, v in source.items() if k != 'chunks'}


def create_source(data, store, ai):
    def text(key, limit, required=True):
        value = data.get(key, '')
        if not isinstance(value, str) or len(value) > limit or (required and not value.strip()):
            raise ValueError(f'{key} must be text, up to {limit} characters' + (' and not empty.' if required else '.'))
        return value.strip()
    target_customer = data.get('customer', store.customer)
    if target_customer != store.customer and store.customer not in ('My Workspace', 'All') and target_customer != 'General':
        raise ValueError('Customer scope is not authorised.')
    if target_customer == 'General' and store.customer not in ('My Workspace', 'All') and data.get('generalized_confirmed') is not True:
        raise ValueError('Confirm that customer names, identifiers, secrets and account-specific details have been removed before sharing this source as General knowledge.')
    content = text('content', MAX_TEXT)
    source_type = data.get('source_type', 'Document')
    if source_type not in ('Document', 'Historical ticket', 'Internal KB', 'SharePoint export', 'Confluence export', 'SAP reference'):
        raise ValueError('Invalid source type.')
    stamp = now()
    source_updated_at = text('source_updated_at', 50, False)
    if source_updated_at:
        try:
            parsed = datetime.fromisoformat(source_updated_at.replace('Z', '+00:00'))
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=timezone.utc)
            source_updated_at = parsed.astimezone(timezone.utc).isoformat()
        except ValueError:
            raise ValueError('Original source update date must be ISO 8601.') from None
    conflict_text = data.get('conflicts_with', '')
    if isinstance(conflict_text, str):
        conflicts = [value.strip() for value in conflict_text.split(',') if value.strip()]
    elif isinstance(conflict_text, list) and all(isinstance(value, str) for value in conflict_text):
        conflicts = [value.strip() for value in conflict_text if value.strip()]
    else:
        raise ValueError('conflicts_with must be a comma-separated list of source IDs.')
    if len(conflicts) > 20 or any(len(value) > 100 for value in conflicts):
        raise ValueError('Provide at most 20 conflict source IDs, each under 100 characters.')
    if any(not store.get(value, 'source') for value in conflicts):
        raise ValueError('Each conflicting source must exist in this customer workspace or be General knowledge.')
    solution_kind = data.get('solution_kind', 'Unknown')
    if solution_kind not in ('Unknown', 'Likely solution', 'Workaround'):
        raise ValueError('Invalid solution kind.')
    source = dict(id='SRC-' + uuid.uuid4().hex[:12].upper(), title=text('title', 200),
        content=content, application=text('application', 100), version=text('version', 100, False),
        platform=text('platform', 150, False), module=text('module', 100, False),
        business_process=text('business_process', 200, False), customisation=text('customisation', 1000, False),
        root_cause=text('root_cause', 3000, False), solution=text('solution', 6000, False),
        solution_kind=solution_kind, source_updated_at=source_updated_at,
        error_code=text('error_code', 100, False), customer=target_customer, source_type=source_type,
        generalization_confirmed=target_customer == 'General' and data.get('generalized_confirmed') is True,
        filename=text('filename', 255, False), source_reference=text('source_reference', 1000, False),
        content_hash=hashlib.sha256(content.encode()).hexdigest(), validation='Unverified',
        validation_evidence={}, steps=[], conflicts_with=list(dict.fromkeys(conflicts)),
        created_at=stamp, updated_at=stamp, index_status='Pending', index_error='', chunk_count=0,
        chunks=[], embedding_model='', revision=0)
    store.save(source, 'source')
    return index_source(source, store, ai)
