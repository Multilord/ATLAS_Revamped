import base64
import io
from unittest.mock import patch
from docx import Document
from test_system import StoreCase
from ingestion import extract_upload, chunks, create_source, index_source
from storage import ConflictError
import server


class WorkflowTests(StoreCase):
    def update(self, ticket, **data):
        return server.update_investigation(ticket, 'workspace', dict(revision=ticket['revision'], **data))

    def test_requirements_and_tests_gate_resolution(self):
        t = self.ticket()
        self.update(t, kind='requirement', title='Vendor fallback?', status='Needs confirmation')
        outcome = dict(status='Resolved', notes='Done', tested=True, customer_confirmed=True, alternative_solution='Reviewed fix')
        with self.assertRaisesRegex(ValueError, 'requirements'):
            server.update_investigation(t, 'outcome', outcome)
        r = t['workflow']['requirements'][0]
        self.update(t, kind='requirement', id=r['id'], title=r['title'], status='Confirmed', evidence='Customer confirmed the rule.')
        self.update(t, kind='scenario', title='Stock shortage', expected='Show shortage', status='Not run')
        with self.assertRaisesRegex(ValueError, 'scenarios'):
            server.update_investigation(t, 'outcome', outcome)
        scenario = t['workflow']['scenarios'][0]
        self.update(t, kind='scenario', id=scenario['id'], title=scenario['title'], expected='Show shortage', status='Passed', actual='Shortage displayed', tester='Reviewer', environment='QA', evidence='Test record 1')
        server.update_investigation(t, 'outcome', outcome)
        self.ai.investigate.assert_not_called()
        self.ai.embed.assert_not_called()

    def test_stage_evidence_and_stale_revision(self):
        t = self.ticket('Change request')
        with self.assertRaises(ValueError):
            self.update(t, kind='stage', name='Approval', status='Completed', owner='Reviewer')
        self.update(t, kind='stage', name='Approval', status='Completed', owner='Reviewer', evidence='Approval record')
        with self.assertRaises(ConflictError):
            server.update_investigation(t, 'workspace', dict(revision=-1, kind='context'))
        with self.assertRaisesRegex(ValueError, 'lifecycle'):
            server.update_investigation(t, 'outcome', dict(status='Resolved', notes='Done'))

    def test_multiple_citations_quote_and_customer_validation(self):
        t=self.ticket()
        s=self.source()
        citation=dict(source_id=s['id'], quote=s['content'], section='Procedure')
        self.update(t, kind='proposal', category='Functional', title='Fix', detail='Review configuration', citations=[citation])
        with self.assertRaisesRegex(ValueError, 'quotation'):
            self.update(t, kind='proposal', category='Technical', title='Fix', detail='Draft', citations=[dict(citation,quote='Fabricated quotation')])
        self.store.db.sources.update_one({'id':s['id']},{'$set':{'customer':'Other Customer'}})
        with self.assertRaisesRegex(ValueError, 'scope'):
            self.update(t, kind='proposal', category='Technical', title='Fix', detail='Draft', citations=[citation])

    def test_local_intake_and_deferred_indexing(self):
        with patch.object(server.ai, 'extract', side_effect=AssertionError('No AI allowed')):
            parsed=server.parse_intake({'content':'Batch selection needs clarification'})
        self.assertIn('Batch selection',parsed['description'])
        source=create_source(dict(title='Procedure',content='Vendor rules',application='SAP',index_now=False),self.store,self.ai)
        self.assertEqual(source['index_status'],'Pending')
        self.ai.embed.assert_not_called()
        indexed=index_source(self.store.get(source['id'],'source',include_chunks=True),self.store,self.ai)
        index_source(self.store.get(indexed['id'],'source',include_chunks=True),self.store,self.ai)
        self.assertEqual(self.ai.embed.call_count,1)

    def test_docx_order_and_section_chunks(self):
        doc=Document()
        doc.add_heading('Functional solution',1)
        doc.add_paragraph('Before table')
        doc.add_table(rows=1,cols=1).cell(0,0).text='Vendor rule'
        doc.add_heading('Technical solution',1)
        doc.add_paragraph('After table')
        output=io.BytesIO();doc.save(output)
        content=extract_upload(dict(filename='example.docx',file_base64=base64.b64encode(output.getvalue()).decode()))['content']
        self.assertLess(content.index('Before table'),content.index('Vendor rule'))
        self.assertLess(content.index('Vendor rule'),content.index('Technical solution'))
        pieces=chunks(content)
        self.assertEqual(len(pieces),2)
        self.assertNotIn('Technical solution',pieces[0]['text'])

    def test_review_is_recorded_at_specific_revision(self):
        t=self.ticket()
        t['analysis']={'summary':'Existing assessment'}
        revision=t['revision']
        self.update(t, kind='evaluation', title='Scenario review', reviewer='Consultant', evidence='Reviewed expected results', summary_complete='Yes', claims_supported='No', scenarios_covered='Not assessed')
        self.assertEqual(t['workflow']['evaluations'][0]['ticket_revision'],revision)
        self.assertEqual(t['analysis']['summary'],'Existing assessment')
        self.ai.embed.assert_not_called()
        self.ai.investigate.assert_not_called()
