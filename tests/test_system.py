"""Isolated test fixtures only. No test documents are loaded by the application."""
import base64
import copy
import io
import json
import os
import threading
import tempfile
import unittest
import uuid
import zipfile
from dataclasses import replace
from datetime import datetime, timezone
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from unittest.mock import Mock, patch

import httpx
import mongomock
from google import genai
from google.genai import types
from openpyxl import Workbook
from docx import Document
from pypdf import PdfWriter
import server
from config import Settings
from storage import MongoStore, ConflictError
from ai import AIService, AIError, GeminiQuota, InvestigationResponse
from engine import retrieve, rank_candidates, conflicting_sources
from ingestion import create_source, extract_upload, chunks, index_source


def config(**changes):
    return replace(Settings(), customer='Test Customer', database='atlas_test_unit', api_key='', **changes)


def response(recommendations=None):
    return InvestigationResponse(classification='Unknown', summary='Review the supplied evidence.', recommendations=recommendations or [], next_steps=['Confirm the release and reproduce the issue in quality.'])


def source_data():
    return dict(title='Test maintenance procedure', content='Inspect the connection configuration and validate the correction in quality.', application='Test App', version='1.0', source_type='Internal KB')


class StoreCase(unittest.TestCase):
    def setUp(self):
        self.cfg = config()
        self.store = MongoStore(self.cfg, client=mongomock.MongoClient())
        self.store.initialize()
        self.ai = Mock()
        self.ai.config = self.cfg
        self.ai.embed.side_effect = lambda texts: [[1., 0., 0.] for _ in texts]
        self.ai.investigate.return_value = response()
        self.old_store, self.old_ai, self.old_customer = server.store, server.ai, server.CUSTOMER
        server.store, server.ai, server.CUSTOMER = self.store, self.ai, self.cfg.customer

    def tearDown(self):
        server.store, server.ai, server.CUSTOMER = self.old_store, self.old_ai, self.old_customer
        self.store.close()

    def ticket(self, kind='Incident'):
        return server.create_investigation(dict(title='Connection problem',description='Inspect connection configuration',application='Test App',version='1.0',type=kind,priority='Medium'))

    def source(self):
        return create_source(source_data(),self.store,self.ai)

    def recommendation(self, source):
        return dict(source_id=source['id'],chunk_id='C0001',quote=source['content'],title='Validate the connection',reason='The supplied procedure describes the symptom.',steps=['Inspect the connection configuration in quality.'],warnings=[],confidence='High',kind='Likely solution')


class DatabaseTests(StoreCase):
    def test_database_starts_empty_and_initialize_never_seeds(self):
        self.assertEqual(self.store.records('investigation'),[])
        self.assertEqual(self.store.records('source'),[])
        self.store.initialize()
        self.assertEqual(self.store.records('source'),[])

    def test_scoped_reads_and_writes(self):
        self.store.db.sources.insert_one(dict(id='private',customer='Other Customer',content='secret',index_status='Ready',embedding_model=self.cfg.embedding_model,chunks=[{'id':'C1','text':'secret','embedding':[1,0,0]}]))
        self.store.db.sources.insert_one(dict(id='legacy-general',customer='General',content='possibly unreviewed',index_status='Ready',embedding_model=self.cfg.embedding_model,chunks=[{'id':'C2','text':'possibly unreviewed','embedding':[1,0,0]}]))
        self.assertIsNone(self.store.get('private','source'))
        self.assertIsNone(self.store.get('legacy-general','source'))
        self.assertEqual(list(self.store.indexed_chunks(self.cfg.embedding_model)),[])
        with self.assertRaises(ValueError):
            self.store.save(dict(id='bad',customer='Other Customer'))

    def test_concurrent_edits_conflict_instead_of_overwriting(self):
        ticket=self.ticket()
        left=self.store.get(ticket['id'],'investigation')
        right=copy.deepcopy(left)
        left['status']='In progress'
        self.store.save(left)
        with self.assertRaises(ConflictError):
            self.store.save(right)

    def test_vectors_never_sent_to_ui(self):
        source=self.source()
        self.assertEqual(source['index_status'],'Ready')
        self.assertNotIn('chunks',self.store.records('source')[0])
        self.assertTrue(self.store.get(source['id'],'source',True)['chunks'])

    def test_duplicate_documents_rejected(self):
        self.source()
        with self.assertRaises(ConflictError):
            self.source()

    def test_model_change_requires_reindex(self):
        self.source()
        self.assertEqual(list(self.store.indexed_chunks('different-model')),[])

    def test_general_sources_require_explicit_generalization(self):
        data=source_data() | dict(customer='General')
        with self.assertRaisesRegex(ValueError,'removed before sharing'):
            create_source(data,self.store,self.ai)
        data['generalized_confirmed']=True
        source=create_source(data,self.store,self.ai)
        self.assertEqual(source['customer'],'General')
        self.assertTrue(source['generalization_confirmed'])

    def test_outcome_signals_require_validation_not_just_resolved_status(self):
        source=self.source()
        self.store.db.investigations.insert_many([
            dict(id='good',customer=self.cfg.customer,selected_source=source['id'],type='Incident',outcome=dict(status='Resolved',tested=True,customer_confirmed=True)),
            dict(id='closed',customer=self.cfg.customer,selected_source=source['id'],type='Incident',outcome=dict(status='Closed - other',tested=True,customer_confirmed=True)),
            dict(id='bad',customer=self.cfg.customer,selected_source=source['id'],type='Incident',outcome=dict(status='Unresolved')),
        ])
        self.assertEqual(self.store.outcome_signals([source['id']])[source['id']],
            dict(validated_successes=1,unsuccessful=1,workarounds=0,partly_useful=0))


class AIGroundingTests(StoreCase):
    def test_analysis_uses_actual_ai_result_and_quote(self):
        source=self.source()
        self.ai.investigate.return_value=response([self.recommendation(source)])
        result=retrieve(self.ticket(),self.store,self.ai)
        self.assertEqual(result['results'][0]['quote'],source['content'])
        self.assertEqual(result['results'][0]['confidence'],'Medium')
        self.assertEqual(result['results'][0]['label'],'Likely solution')
        self.assertEqual(result['ai_solutioning']['provenance'],'AI-generated; not consultant-verified')
        self.assertEqual(result['ai_solutioning']['functional_assessment'],'')

    def test_fabricated_citation_rejected(self):
        source=self.source()
        rec=self.recommendation(source)
        rec['source_id']='invented'
        self.ai.investigate.return_value=response([rec])
        with self.assertRaises(AIError):
            retrieve(self.ticket(),self.store,self.ai)

    def test_fabricated_quote_rejected(self):
        source=self.source()
        rec=self.recommendation(source)
        rec['quote']='A fabricated supporting excerpt.'
        self.ai.investigate.return_value=response([rec])
        with self.assertRaises(AIError):
            retrieve(self.ticket(),self.store,self.ai)

    def test_empty_corpus_returns_no_fake_recommendations(self):
        result=retrieve(self.ticket(),self.store,self.ai)
        self.assertEqual(result['state'],'Insufficient evidence')
        self.assertEqual(result['results'],[])
        self.assertEqual(self.ai.investigate.call_args.args[1],[])

    def test_version_mismatch_cannot_be_high_confidence(self):
        source=self.source()
        self.ai.investigate.return_value=response([self.recommendation(source)])
        ticket=self.ticket()
        ticket['version']='2.0'
        result=retrieve(ticket,self.store,self.ai)
        self.assertEqual(result['results'][0]['confidence'],'Low')
        self.assertEqual(result['state'],'Insufficient evidence')

    def test_contextual_metadata_boosts_matching_module_and_platform(self):
        matching=create_source(source_data() | dict(title='FI matching source',module='FI',platform='S/4HANA',business_process='invoice posting'),self.store,self.ai)
        other=create_source(source_data() | dict(title='Other module source',content='Different procedure for connection settings and validation.',module='MM',platform='ECC'),self.store,self.ai)
        ticket=self.ticket()
        ticket.update(module='FI',platform='S/4HANA',business_process='invoice posting')
        ranked,_=rank_candidates(ticket,self.store,self.ai)
        scores={row['source']['id']:row['score'] for row in ranked}
        self.assertGreater(scores[matching['id']],scores[other['id']])

    def test_conflicting_sources_report_both_ids(self):
        left=dict(source=dict(id='A',error_code='E-1',application='SAP FI',version='v1',solution='Open period'))
        right=dict(source=dict(id='B',error_code='E-1',application='SAP FI',version='v1',solution='Do not open period'))
        conflicts=conflicting_sources([left,right])
        self.assertEqual(conflicts['A'],{'B'})
        self.assertEqual(conflicts['B'],{'A'})

    def test_indexing_failure_preserves_real_document_for_retry(self):
        self.ai.embed.side_effect=AIError('API key missing')
        source=self.source()
        self.assertEqual(source['index_status'],'Error')
        self.assertEqual(self.store.get(source['id'],'source')['content'],source_data()['content'])
        self.assertEqual(list(self.store.indexed_chunks(self.cfg.embedding_model)),[])
        self.ai.embed.side_effect=lambda texts:[[1.,0.,0.] for _ in texts]
        updated=index_source(self.store.get(source['id'],'source',True),self.store,self.ai)
        self.assertEqual(updated['index_status'],'Ready')

    def test_failed_ai_analysis_does_not_save_fake_success(self):
        ticket=self.ticket()
        self.ai.investigate.side_effect=AIError('Provider failed')
        with self.assertRaises(AIError):
            server.update_investigation(ticket,'analyze',{})
        self.assertIsNone(self.store.get(ticket['id'],'investigation')['analysis'])


class WorkflowTests(StoreCase):
    def test_full_workflow_persistence(self):
        source=self.source()
        self.ai.investigate.return_value=response([self.recommendation(source)])
        ticket=server.update_investigation(self.ticket(),'analyze',{})
        ticket=server.update_investigation(ticket,'decision',dict(source_id=source['id'],decision='Accepted'))
        ticket=server.update_investigation(ticket,'outcome',dict(status='Resolved',notes='Validated against quality test record.',tested=True,customer_confirmed=True))
        stored=self.store.get(ticket['id'],'investigation')
        self.assertEqual(stored['status'],'Awaiting validation')
        self.assertEqual(stored['outcome']['status'],'Resolved')
        self.assertEqual(stored['lifecycle_stage'],'Ticket Creation')
        self.assertEqual(len(stored['timeline']),4)

    def test_invalid_resolution_validation(self):
        for flags in [dict(tested=False,customer_confirmed=True),dict(tested='yes',customer_confirmed=True)]:
            with self.subTest(flags=flags),self.assertRaises(ValueError):
                server.update_investigation(self.ticket(),'outcome',dict(status='Resolved',notes='Reviewed',alternative_solution='Procedure',**flags))

    def test_change_request_requires_change_stages(self):
        with self.assertRaises(ValueError):
            server.update_investigation(self.ticket('Change request'),'outcome',dict(status='Resolved',notes='Reviewed',tested=True,customer_confirmed=True,alternative_solution='Procedure'))

    def advance(self,item,stage,**extra):
        return server.transition_lifecycle(item,dict(stage=stage,notes=f'Recorded {stage} action.',**extra))

    def test_incident_and_service_request_follow_standard_lifecycle(self):
        for kind in ('Incident','Service request'):
            with self.subTest(kind=kind):
                item=self.ticket(kind)
                self.assertEqual(item['lifecycle_stage'],'Ticket Creation')
                for stage in ('Ticket Acceptance','Acknowledgment','Clarification'):
                    item=self.advance(item,stage)
                with self.assertRaisesRegex(ValueError,'proposed solution'):
                    self.advance(item,'Solutioning')
                item=server.update_solutioning(item,dict(section='functional',fields={'proposed_solution':'Follow the validated customer process.'}))
                item=self.advance(item,'Solutioning')
                item=server.update_investigation(item,'outcome',dict(status='Worked',notes='Validated by customer.',alternative_solution='Follow the validated customer process.',tested=True,customer_confirmed=True))
                item=self.advance(item,'Closure')
                saved=self.store.get(item['id'],'investigation')
                self.assertEqual(saved['type'],kind)
                self.assertEqual(saved['lifecycle_stage'],'Closure')
                self.assertEqual([e['stage'] for e in saved['lifecycle_history']],['Ticket Creation','Ticket Acceptance','Acknowledgment','Clarification','Solutioning','Closure'])
                self.assertEqual(saved['outcome']['status'],'Worked')
                self.assertEqual(saved['status'],'Closed')

    def test_invalid_lifecycle_skip_rejected(self):
        item=self.ticket()
        with self.assertRaisesRegex(ValueError,'exactly one stage'):
            self.advance(item,'Clarification')

    def test_change_request_requires_approval_development_deployment_verification_then_closure(self):
        item=self.ticket('Change request')
        for stage in ('Ticket Acceptance','Acknowledgment','Clarification'):
            item=self.advance(item,stage)
        item=server.update_solutioning(item,dict(section='functional',fields={'proposed_solution':'Add the requested business behavior.'}))
        item=self.advance(item,'Solutioning')
        item=self.advance(item,'Approval',approval_status='Pending')
        with self.assertRaisesRegex(ValueError,'cannot start until'):
            self.advance(item,'Development')
        item=server.record_approval(item,dict(status='Approved',notes='Customer approval reference CR-7.'))
        item=self.advance(item,'Development')
        item=server.update_solutioning(item,dict(section='technical',fields={'implementation_approach':'Implement the approved behavior and deploy via change controls.'}))
        item=self.advance(item,'Deployment')
        item=self.advance(item,'Verification')
        item=server.record_verification(item,dict(result='Passed',notes='Verification passed in the quality environment.'))
        item=server.update_investigation(item,'outcome',dict(status='Worked',notes='Customer confirmed after verification.',alternative_solution='Approved implementation',tested=True,customer_confirmed=True,approved=True,developed=True,deployed=True,verified=True))
        item=self.advance(item,'Closure')
        saved=self.store.get(item['id'],'investigation')
        stages=[entry['stage'] for entry in saved['lifecycle_history']]
        self.assertEqual(stages,['Ticket Creation','Ticket Acceptance','Acknowledgment','Clarification','Solutioning','Approval','Development','Deployment','Verification','Closure'])
        self.assertEqual(saved['verification']['result'],'Passed')
        self.assertEqual(saved['approval_decision']['status'],'Approved')

    def test_change_cannot_close_before_verification_passes(self):
        item=self.ticket('Change request')
        for stage in ('Ticket Acceptance','Acknowledgment','Clarification'):
            item=self.advance(item,stage)
        item=server.update_solutioning(item,dict(section='technical',fields={'implementation_approach':'Implement after customer approval.'}))
        item=self.advance(item,'Solutioning')
        item=self.advance(item,'Approval',approval_status='Approved')
        item=self.advance(item,'Development')
        item=self.advance(item,'Deployment')
        item=self.advance(item,'Verification',verification_result='Failed')
        item=server.update_investigation(item,'outcome',dict(status='Partially worked',notes='Verification did not pass.',alternative_solution='Pending fix',tested=True))
        with self.assertRaisesRegex(ValueError,'verification has passed'):
            self.advance(item,'Closure')

    def test_verification_is_separate_repeatable_and_persisted(self):
        item=self.ticket('Change request')
        for stage in ('Ticket Acceptance','Acknowledgment','Clarification'):
            item=self.advance(item,stage)
        item=server.update_solutioning(item,dict(section='technical',fields={'implementation_approach':'Approved implementation proposal.'}))
        item=self.advance(item,'Solutioning')
        item=self.advance(item,'Approval',approval_status='Approved')
        item=self.advance(item,'Development')
        item=self.advance(item,'Deployment')
        item=self.advance(item,'Verification')
        self.assertEqual(item['verification']['result'],'Pending')
        item=server.record_verification(item,dict(result='Failed',notes='Expected output did not match.'))
        self.assertEqual(item['lifecycle_stage'],'Verification')
        item=server.record_verification(item,dict(result='Passed',notes='Retest passed; evidence TEST-9.'))
        saved=self.store.get(item['id'],'investigation')
        self.assertEqual(saved['verification']['result'],'Passed')
        self.assertEqual([entry['result'] for entry in saved['verification_history']],['Pending','Failed','Passed'])
        self.assertEqual(len([e for e in saved['timeline'] if e['kind']=='Verification']),3)

    def test_legacy_terminal_ticket_does_not_get_invented_lifecycle_history(self):
        item=self.ticket()
        item['status']='Resolved'
        item.pop('lifecycle_history')
        item.pop('lifecycle_history_tracked')
        item.pop('lifecycle_stage')
        view=server.lifecycle_view(item)
        self.assertFalse(view['history_tracked'])
        self.assertIsNone(view['next_stage'])
        self.assertEqual(view['current_stage'],'Lifecycle not tracked for this imported record')
        with self.assertRaisesRegex(ValueError,'already terminal'):
            self.advance(item,'Ticket Acceptance')

    def test_functional_and_technical_solutioning_are_separate_and_persist(self):
        item=self.ticket()
        item=server.update_solutioning(item,dict(section='functional',fields={'expected_behaviour':'Customer expects one approval.'}))
        item=server.update_solutioning(item,dict(section='technical',fields={'risks':'Interface compatibility must be checked.'}))
        saved=self.store.get(item['id'],'investigation')
        self.assertEqual(saved['solutioning']['functional']['customer_problem'],saved['description'])
        self.assertEqual(saved['solutioning']['functional']['expected_behaviour'],'Customer expects one approval.')
        self.assertEqual(saved['solutioning']['technical']['risks'],'Interface compatibility must be checked.')
        self.assertEqual(len(saved['solutioning_history']),2)
        self.assertEqual([event['source'] for event in saved['timeline'] if event['source'] in ('Functional solutioning','Technical solutioning')],['Functional solutioning','Technical solutioning'])

    def test_outcome_updates_preserve_feedback_and_history(self):
        item=self.ticket()
        original_status=item['status']
        item=server.update_investigation(item,'outcome',dict(status='Partially worked',notes='First result.',usefulness='Partly useful'))
        self.assertEqual(item['status'],original_status)
        self.assertEqual(item['lifecycle_stage'],'Ticket Creation')
        item=server.update_investigation(item,'outcome',dict(status='Did not work',notes='Retest failed.',usefulness='Not useful'))
        saved=self.store.get(item['id'],'investigation')
        self.assertEqual(saved['outcome']['status'],'Did not work')
        self.assertEqual(saved['outcome_history'][0]['status'],'Partially worked')
        self.assertEqual(len([event for event in saved['timeline'] if event['kind']=='Outcome']),2)

    def test_new_evidence_invalidates_recommendations(self):
        ticket=server.update_investigation(self.ticket(),'analyze',{})
        ticket=server.update_investigation(ticket,'note',dict(kind='Test',text='A different result was observed.'))
        self.assertIsNone(ticket['analysis'])

    def test_search_evidence_does_not_advance_new_ticket_status_or_lifecycle(self):
        ticket=self.ticket()
        ticket=server.update_investigation(ticket,'analyze',{})
        self.assertEqual(ticket['status'],'New')
        self.assertEqual(ticket['lifecycle_stage'],'Ticket Creation')
        self.assertIsNotNone(ticket['analysis'])

    def test_search_evidence_is_available_after_outcome_and_closure_without_reopening(self):
        ticket=self.ticket()
        ticket['status']='Closed'
        ticket['lifecycle_stage']='Closure'
        ticket['lifecycle_history'].append(dict(stage='Closure',at=server.now(),author='Tester',notes='Closed.'))
        ticket['outcome']=dict(status='Worked',notes='Previously validated.',recorded_at=server.now(),selected_source='SRC-KEEP',usefulness='Useful',tested=True,customer_confirmed=True)
        ticket['selected_source']='SRC-KEEP'
        ticket['decisions']={'SRC-KEEP':'Accepted'}
        self.store.save(ticket)
        refreshed=self.store.get(ticket['id'],'investigation')
        analyzed=server.update_investigation(refreshed,'analyze',{})
        persisted=self.store.get(ticket['id'],'investigation')
        self.assertIsNotNone(persisted['analysis'])
        self.assertEqual(persisted['status'],'Closed')
        self.assertEqual(persisted['lifecycle_stage'],'Closure')
        self.assertEqual(persisted['outcome']['notes'],'Previously validated.')
        self.assertEqual(persisted['selected_source'],'SRC-KEEP')
        self.assertEqual(persisted['decisions'],{'SRC-KEEP':'Accepted'})
        self.assertEqual(analyzed['status'],'Closed')

    def test_imported_chronology_is_preserved(self):
        data=dict(title='Actual ticket',description='Details',timeline=[dict(at='2026-10-14T10:00:00+08:00',text='Later'),dict(at='2026-10-14T00:00:00Z',text='Earlier')])
        result=server.parse_intake(dict(content=json.dumps(data)),use_ai=True)
        self.assertEqual([e['text'] for e in result['timeline']],['Earlier','Later'])
        self.ai.extract.assert_not_called()

    def test_imported_timeline_event_kinds_are_preserved(self):
        data=dict(description='Problem',timeline=[dict(at='2026-10-14T00:00:00Z',text='Customer confirmed steps',kind='Clarification',author='Customer',source='mail-1')])
        event=server.parse_intake(dict(content=json.dumps(data)))['timeline'][0]
        self.assertEqual((event['kind'],event['author'],event['source']),('Clarification','Customer','mail-1'))

    def test_unstructured_intake_calls_ai(self):
        self.ai.extract.return_value=dict(title='Extracted title',description='Reported issue',timeline=[])
        parsed=server.parse_intake(dict(content='Unstructured original ticket'),use_ai=True)
        self.ai.extract.assert_called_once_with('Unstructured original ticket')
        self.assertEqual(parsed['title'],'Extracted title')

    def test_other_customer_intake_rejected(self):
        with self.assertRaises(ValueError):
            server.parse_intake(dict(content=json.dumps({'customer':'Other Customer'})))


class FileTests(unittest.TestCase):
    def upload(self,name,content):
        return extract_upload(dict(filename=name,file_base64=base64.b64encode(content).decode()))

    def test_utf8_and_docx_extraction(self):
        self.assertEqual(self.upload('ticket.txt','Invoice café'.encode())['content'],'Invoice café')
        doc=Document()
        doc.add_paragraph('Observed application error')
        data=io.BytesIO()
        doc.save(data)
        self.assertIn('Observed application error',self.upload('source.docx',data.getvalue())['content'])

    def test_scanned_or_blank_pdf_explicitly_rejected(self):
        writer=PdfWriter()
        writer.add_blank_page(width=200,height=200)
        data=io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(ValueError,'OCR'):
            self.upload('blank.pdf',data.getvalue())

    def test_spreadsheet_and_csv_extraction(self):
        workbook=Workbook()
        workbook.active.title='Tickets'
        workbook.active.append(['ID','Summary'])
        workbook.active.append(['INC-42','Posting period closed'])
        data=io.BytesIO()
        workbook.save(data)
        workbook.close()
        extracted=self.upload('tickets.xlsx',data.getvalue())['content']
        self.assertIn('[Sheet: Tickets]',extracted)
        self.assertIn('INC-42 | Posting period closed',extracted)
        self.assertEqual(self.upload('tickets.csv',b'ID,Summary\r\nINC-42,Closed')['content'],'ID,Summary\r\nINC-42,Closed')

    def test_html_and_powerpoint_extraction(self):
        html=self.upload('guide.html',b'<h1>Posting guide</h1><script>private_code()</script><p>Check the period.</p>')['content']
        self.assertIn('Posting guide',html)
        self.assertIn('Check the period.',html)
        self.assertNotIn('private_code',html)
        pptx=io.BytesIO()
        with zipfile.ZipFile(pptx,'w') as archive:
            archive.writestr('ppt/slides/slide1.xml','<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Validate in quality</a:t></p:sld>')
        self.assertIn('Validate in quality',self.upload('guide.pptx',pptx.getvalue())['content'])

    def test_bad_and_unsupported_uploads(self):
        for name,data in [('bad.json',b'{broken'),('script.exe',b'content'),('empty.txt',b'')]:
            with self.subTest(name=name),self.assertRaises(ValueError):
                self.upload(name,data)

    def test_chunk_coverage_and_boundaries(self):
        text='word ' * 1400
        parts=chunks(text)
        self.assertGreater(len(parts),1)
        self.assertTrue(all(p['text'] in text and len(p['text'])<=2200 for p in parts))
        self.assertTrue(text.rstrip().endswith(parts[-1]['text']))


class ProviderTests(unittest.TestCase):
    def test_missing_key_is_explicit(self):
        with self.assertRaisesRegex(AIError,'GEMINI_API_KEY'):
            AIService(config()).embed(['text'])

    def client(self, transport):
        return genai.Client(api_key='test-key-not-real', vertexai=False,
            http_options=types.HttpOptions(httpx_client=httpx.Client(transport=httpx.MockTransport(transport)),
                retry_options=types.HttpRetryOptions(attempts=1)))

    def generation(self, text=None, finish='STOP'):
        return {'candidates':[{'content':{'role':'model','parts':[{'text':text if text is not None else response().model_dump_json()}]},'finishReason':finish}]}

    def test_generation_and_embeddings_wire_contract(self):
        requests=[]
        def transport(request):
            body=json.loads(request.content)
            requests.append((request.url.path,body))
            if 'embedContent' in request.url.path or 'batchEmbedContents' in request.url.path:
                return httpx.Response(200,json={'embeddings':[{'values':[1.,0.,0.]}]})
            return httpx.Response(200,json=self.generation())
        cfg=replace(config(),api_key='test-key-not-real',model='gemini-2.5-flash')
        with self.client(transport) as client:
            ai=AIService(cfg,client)
            self.assertEqual(ai.embed(['ticket']),[[1.,0.,0.]])
            self.assertEqual(ai.investigate({'title':'Ticket'},[]).recommendations,[])
        self.assertIn(cfg.embedding_model,requests[0][0])
        self.assertIn(cfg.model,requests[1][0])
        generation=requests[1][1]['generationConfig']
        self.assertEqual(generation['responseMimeType'],'application/json')
        self.assertEqual(generation['thinkingConfig'].get('thinkingBudget',generation['thinkingConfig'].get('thinking_budget')),0)
        self.assertNotIn('responseJsonSchema',generation)
        instruction=requests[1][1]['systemInstruction']['parts'][0]['text']
        self.assertIn('"recommendations"',instruction)
        self.assertIn('untrusted DATA',instruction)

    def test_auth_error_never_exposes_key_or_provider_body(self):
        cfg=replace(config(),api_key='secret-test-key')
        with self.client(lambda r:httpx.Response(403,json={'error':{'code':403,'message':'secret-test-key','status':'PERMISSION_DENIED'}})) as client:
            with self.assertRaises(AIError) as caught:
                AIService(cfg,client).embed(['text'])
        self.assertNotIn('secret-test-key',str(caught.exception))
        self.assertIn('GEMINI_API_KEY',str(caught.exception))

    def test_malformed_and_incomplete_generation_rejected(self):
        cfg=replace(config(),api_key='test-key-not-real')
        for payload in [self.generation('{}'),self.generation(finish='MAX_TOKENS'),{'candidates':[]}]:
            with self.subTest(payload=payload),self.client(lambda r:httpx.Response(200,json=payload)) as client:
                with self.assertRaises(AIError):
                    AIService(cfg,client).investigate({'title':'Ticket'},[])

    def test_ticket_extraction_validates_gemini_json(self):
        cfg=replace(config(),api_key='test-key-not-real')
        ticket=dict(title='Reported failure',description='Original symptom',ticket_id='',application='Test App',
            platform='',module='',business_process='',version='1.0',error_code='',customisation='',type='Incident',priority='Medium',timeline=[])
        with self.client(lambda r:httpx.Response(200,json=self.generation(json.dumps(ticket)))) as client:
            self.assertEqual(AIService(cfg,client).extract('Original ticket'),ticket)

    def test_settings_use_gemini_environment_only(self):
        with patch.dict(os.environ,{'GEMINI_API_KEY':'gemini-test-key','GEMINI_MODEL':'custom-gemini',
                'GEMINI_EMBEDDING_MODEL':'custom-embedding','GEMINI_TIMEOUT_SECONDS':'12',
                'OPENAI_API_KEY':'old-provider-key'},clear=True):
            cfg=Settings()
            self.assertEqual(cfg.api_key,'gemini-test-key')
            self.assertEqual((cfg.model,cfg.embedding_model,cfg.ai_timeout),('custom-gemini','custom-embedding',12))
            self.assertNotIn('gemini-test-key',repr(cfg))
        with patch.dict(os.environ,{'OPENAI_API_KEY':'old-provider-key'},clear=True):
            self.assertFalse(Settings().api_key)

    def test_embedding_count_and_nonfinite_values_rejected(self):
        cfg=replace(config(),api_key='test-key-not-real')
        for payload in [{'embeddings':[]},{'embeddings':[{'values':[]}]},{'embeddings':[{'values':['NaN']}]}]:
            with self.subTest(payload=payload),self.client(lambda r:httpx.Response(200,json=payload)) as client:
                with self.assertRaises(AIError):
                    AIService(cfg,client).embed(['ticket'])

    def test_batch_preserves_separate_text_embeddings(self):
        requests=[]
        def transport(request):
            body=json.loads(request.content)
            requests.append(body)
            count=len(body['requests'])
            return httpx.Response(200,json={'embeddings':[{'values':[float(i),1.]} for i in range(count)]})
        cfg=replace(config(),api_key='test-key-not-real')
        with self.client(transport) as client:
            vectors=AIService(cfg,client).embed([f'text {i}' for i in range(17)])
        self.assertEqual(len(vectors),17)
        self.assertEqual([len(r['requests']) for r in requests],[16,1])
        self.assertEqual(requests[0]['requests'][1]['content']['parts'][0]['text'],'text 1')

    def test_quota_and_network_errors_are_actionable(self):
        cfg=replace(config(),api_key='test-key-not-real')
        with self.client(lambda r:httpx.Response(429,json={'error':{'code':429,'message':'private detail','status':'RESOURCE_EXHAUSTED'}})) as client:
            with self.assertRaisesRegex(AIError,'quota'):
                AIService(cfg,client).embed(['ticket'])
        def offline(request):
            raise httpx.ConnectError('private transport detail')
        with self.client(offline) as client:
            with self.assertRaisesRegex(AIError,'network'):
                AIService(cfg,client).embed(['ticket'])

    def test_connection_check_uses_configured_models(self):
        cfg=replace(config(),api_key='test-key-not-real')
        paths=[]
        def transport(request):
            paths.append(request.url.path)
            if 'generateContent' in request.url.path:
                return httpx.Response(200,json=self.generation('{"status":"ok"}'))
            return httpx.Response(200,json={'embeddings':[{'values':[1.,0.,0.]}]})
        with self.client(transport) as client:
            self.assertEqual(AIService(cfg,client).check()['status'],'Connected')
        self.assertIn(cfg.model+':generateContent',paths[0])
        self.assertEqual(len(paths),1)

    def test_json_mode_still_enforces_limits_and_extra_fields(self):
        cfg=replace(config(),api_key='test-key-not-real')
        for payload in [response().model_dump() | {'summary':'x'*3001},
                response().model_dump() | {'unexpected':'value'},
                response().model_dump() | {'next_steps':['step']*9}]:
            with self.subTest(payload=payload),self.client(lambda r:httpx.Response(200,json=self.generation(json.dumps(payload)))) as client:
                with self.assertRaisesRegex(AIError,'schema validation'):
                    AIService(cfg,client).investigate({'title':'Ticket'},[])

    def test_ticket_schema_preserves_title_and_local_timeline_limit(self):
        cfg=replace(config(),api_key='test-key-not-real')
        requests=[]
        ticket=dict(title='Test ticket',description='Symptom',ticket_id='',application='App',platform='',module='',business_process='',version='1',
            error_code='',customisation='',type='Incident',priority='Medium',timeline=[])
        def transport(request):
            requests.append(json.loads(request.content))
            return httpx.Response(200,json=self.generation(json.dumps(ticket)))
        with self.client(transport) as client:
            AIService(cfg,client).extract('Original ticket')
        instruction=requests[0]['systemInstruction']['parts'][0]['text']
        schema=json.loads(instruction.split('matching this schema:\n',1)[1])
        self.assertIn('title',schema['properties'])
        self.assertIn('title',schema['required'])
        self.assertNotIn('maxItems',schema['properties']['timeline'])
        ticket['timeline']=[dict(at='2026-10-14T00:00:00Z',author='Tester',text='Event')]*101
        with self.client(lambda r:httpx.Response(200,json=self.generation(json.dumps(ticket)))) as client:
            with self.assertRaisesRegex(AIError,'schema validation'):
                AIService(cfg,client).extract('Original ticket')

    def test_connection_check_rejects_failed_generation(self):
        cfg=replace(config(),api_key='test-key-not-real')
        with self.client(lambda r:httpx.Response(503,json={'error':{'code':503,'message':'private detail','status':'UNAVAILABLE'}})) as client:
            with self.assertRaisesRegex(AIError,'temporarily unavailable'):
                AIService(cfg,client).check()

    def test_busy_provider_returns_retry_guidance(self):
        cfg=replace(config(),api_key='test-key-not-real')
        with self.client(lambda r:httpx.Response(503,json={'error':{'code':503,'message':'private detail','status':'UNAVAILABLE'}})) as client:
            with self.assertRaisesRegex(AIError,'temporarily unavailable'):
                AIService(cfg,client).investigate({'title':'Ticket'},[])


class LocalQuotaTests(unittest.TestCase):
    def setUp(self):
        self.directory=tempfile.TemporaryDirectory()
        self.now=[datetime(2026,10,14,12,tzinfo=timezone.utc).timestamp()]
        self.path=os.path.join(self.directory.name,'quota.json')
        self.quota=GeminiQuota(5,20,self.path,clock=lambda:self.now[0])

    def tearDown(self):
        self.directory.cleanup()

    def test_minute_limit_and_retry(self):
        for _ in range(5):
            self.quota.acquire()
        status=self.quota.status()
        self.assertEqual((status['minute_remaining'],status['daily_remaining']),(0,15))
        with self.assertRaisesRegex(AIError,'5 requests per minute'):
            self.quota.acquire()
        self.now[0]+=61
        self.quota.acquire()
        self.assertEqual(self.quota.status()['daily_used'],6)

    def test_daily_limit_persists_across_instances(self):
        for _ in range(4):
            for _ in range(5):
                self.quota.acquire()
            self.now[0]+=61
        reopened=GeminiQuota(5,20,self.path,clock=lambda:self.now[0])
        self.assertEqual(reopened.status()['daily_remaining'],0)
        with self.assertRaisesRegex(AIError,'daily limit'):
            reopened.acquire()

    def test_preflight_reserves_multi_request_operation_without_consuming(self):
        for _ in range(4):
            self.quota.acquire()
        with self.assertRaisesRegex(AIError,'needs 2 Gemini request'):
            self.quota.ensure(2)
        self.assertEqual(self.quota.status()['minute_used'],4)

    def test_utc_day_reset(self):
        self.quota.acquire()
        self.now[0]=datetime(2026,10,15,0,0,1,tzinfo=timezone.utc).timestamp()
        self.assertEqual(self.quota.status()['daily_used'],0)

    def test_invalid_limits_rejected(self):
        with self.assertRaises(ValueError):
            GeminiQuota(0,20,self.path)


class HttpTests(StoreCase):
    def setUp(self):
        super().setUp()
        server.ai=AIService(self.cfg)
        self.service=server.ThreadingHTTPServer(('127.0.0.1',0),server.Handler)
        self.thread=threading.Thread(target=self.service.serve_forever,daemon=True)
        self.thread.start()
        self.url=f'http://127.0.0.1:{self.service.server_port}'

    def tearDown(self):
        self.service.shutdown()
        self.service.server_close()
        self.thread.join()
        super().tearDown()

    def test_state_empty_and_no_credentials_exposed(self):
        with urlopen(self.url+'/api/state') as res:
            state=json.load(res)
        self.assertEqual(state['investigations'],[])
        self.assertEqual(state['sources'],[])
        self.assertNotIn('mongo_uri',json.dumps(state))
        self.assertEqual(state['services']['ai'],'API key missing')

    def test_ai_without_key_returns_actionable_failure(self):
        request=Request(self.url+'/api/ai/check',data=b'{}',headers={'Content-Type':'application/json'})
        with self.assertRaises(HTTPError) as caught:
            urlopen(request)
        self.assertEqual(caught.exception.code,502)
        self.assertIn('GEMINI_API_KEY',caught.exception.read().decode())

    def test_cross_origin_rejected(self):
        request=Request(self.url+'/api/sources',data=b'{}',headers={'Content-Type':'application/json','Origin':'https://example.com'})
        with self.assertRaises(HTTPError) as caught:
            urlopen(request)
        self.assertEqual(caught.exception.code,403)

    def test_lifecycle_endpoint_persists_stage_and_exposes_it_in_state(self):
        item=self.ticket()
        data=json.dumps(dict(stage='Ticket Acceptance',notes='Accepted by the support consultant.')).encode()
        request=Request(self.url+f"/api/investigations/{item['id']}/lifecycle",data=data,headers={'Content-Type':'application/json'})
        with urlopen(request) as response:
            changed=json.load(response)
        self.assertEqual(changed['lifecycle_stage'],'Ticket Acceptance')
        with urlopen(self.url+'/api/state') as response:
            state=json.load(response)
        persisted=next(record for record in state['investigations'] if record['id']==item['id'])
        self.assertEqual(persisted['lifecycle']['current_stage'],'Ticket Acceptance')
        self.assertEqual([entry['stage'] for entry in persisted['lifecycle_history']],['Ticket Creation','Ticket Acceptance'])

    def test_solutioning_endpoint_persists_both_sections(self):
        item=self.ticket()
        for section,fields in [('functional',{'expected_behaviour':'Customer expects a guided approval.'}),('technical',{'risks':'Integration compatibility requires testing.'})]:
            body=json.dumps(dict(section=section,fields=fields)).encode()
            request=Request(self.url+f"/api/investigations/{item['id']}/solutioning",data=body,headers={'Content-Type':'application/json'})
            with urlopen(request) as response:
                saved=json.load(response)
        with urlopen(self.url+'/api/state') as response:
            state=json.load(response)
        persisted=next(record for record in state['investigations'] if record['id']==item['id'])
        self.assertEqual(persisted['solutioning']['functional']['expected_behaviour'],'Customer expects a guided approval.')
        self.assertEqual(persisted['solutioning']['technical']['risks'],'Integration compatibility requires testing.')
        self.assertEqual(saved['id'],item['id'])


@unittest.skipUnless(os.getenv('RUN_MONGO_TESTS')=='1','Set RUN_MONGO_TESTS=1 to test a real MongoDB instance.')
class RealMongoTests(unittest.TestCase):
    def test_real_persistence_and_empty_initial_state(self):
        # Atlas limits database names to 38 bytes; keep the disposable name
        # comfortably below that while retaining enough random entropy.
        name='atlas_test_'+uuid.uuid4().hex[:20]
        cfg=replace(config(),database=name)
        store=MongoStore(cfg)
        try:
            store.initialize()
            self.assertEqual(store.records('source'),[])
            item=dict(id='integration-record',customer=cfg.customer,updated_at=datetime.now(timezone.utc).isoformat(),timeline=[],status='New')
            store.save(item)
            reopened=MongoStore(cfg)
            try:
                self.assertEqual(reopened.get(item['id'],'investigation')['status'],'New')
            finally:
                reopened.close()
        finally:
            assert name.startswith('atlas_test_') and name!=Settings().database
            store.client.drop_database(name)
            store.close()


if __name__=='__main__':
    unittest.main()
